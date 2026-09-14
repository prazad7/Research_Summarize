import openpyxl
import pytest

from app.db.models import ContentType, Job, JobStatus
from app.ingestion.excel_extractor import ExcelExtractor
from app.ingestion.exceptions import CorruptFileError, NoExtractableTextError


def _make_job(path: str) -> Job:
    return Job(
        id="test-job",
        status=JobStatus.PENDING,
        content_type=ContentType.XLSX,
        source_type="file",
        source_reference="budget.xlsx",
        stored_file_path=path,
    )


def test_excel_extractor_reads_xlsx_rows(tmp_path):
    xlsx_path = tmp_path / "budget.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Q1"
    sheet.append(["Name", "Amount"])
    sheet.append(["Rent", 1200])
    workbook.save(str(xlsx_path))

    extracted = ExcelExtractor().extract(_make_job(str(xlsx_path)))

    assert "Q1" in extracted.text
    assert "Rent" in extracted.text
    assert "1200" in extracted.text
    assert extracted.title == "budget"
    assert extracted.extra_metadata["sheet_count"] == 1


def test_excel_extractor_raises_on_empty_workbook(tmp_path):
    xlsx_path = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(str(xlsx_path))

    with pytest.raises(NoExtractableTextError):
        ExcelExtractor().extract(_make_job(str(xlsx_path)))


def test_excel_extractor_raises_corrupt_file_error_on_non_zip_content(tmp_path):
    """Regression test: a file that isn't a valid zip container at all (a
    legacy binary .xls mislabeled as .xlsx, a truncated download, or an
    unrelated file) used to raise a raw, uncaught `zipfile.BadZipFile`
    instead of a friendly IngestionError -- surfacing as the generic
    "Something went wrong" message instead of a clear one."""
    bad_path = tmp_path / "not_really.xlsx"
    bad_path.write_text("this is not a valid xlsx file")

    with pytest.raises(CorruptFileError):
        ExcelExtractor().extract(_make_job(str(bad_path)))


def test_excel_extractor_reads_legacy_xls_via_xlrd(tmp_path, mocker):
    xls_path = tmp_path / "legacy.xls"
    xls_path.write_bytes(b"not a real xls, just needs to exist on disk")

    fake_sheet = mocker.MagicMock(name="Sheet1", nrows=2)
    fake_sheet.name = "Sheet1"
    fake_sheet.row_values.side_effect = lambda r: {
        0: ["Name", "Amount"],
        1: ["Rent", 1200],
    }[r]
    fake_workbook = mocker.MagicMock(nsheets=1)
    fake_workbook.sheets.return_value = [fake_sheet]

    mocker.patch("xlrd.open_workbook", return_value=fake_workbook)

    extracted = ExcelExtractor().extract(_make_job(str(xls_path)))

    assert "Sheet1" in extracted.text
    assert "Rent" in extracted.text
    assert extracted.extra_metadata["sheet_count"] == 1


def test_excel_extractor_keeps_column_alignment_when_a_middle_cell_is_blank(tmp_path):
    """Regression test for a real reported bug: a blank cell in a middle
    column (e.g. Phone) used to be dropped instead of kept as an empty
    field, shifting every later value (Course, Fee_Paid, City, ...) one
    column to the left for that row -- silently corrupting the row for
    anyone reading it by column position, including the chat LLM, which
    undercounted a "how many from Chennai" query because misaligned rows
    no longer had "Chennai" positioned under the City column."""
    xlsx_path = tmp_path / "students.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Name", "Phone", "Course", "City"])
    sheet.append(["Alice", "9876543210", "python", "Chennai"])
    sheet.append(["Bob", None, "python", "Chennai"])  # blank Phone

    workbook.save(str(xlsx_path))
    extracted = ExcelExtractor().extract(_make_job(str(xlsx_path)))

    lines = extracted.text.splitlines()
    alice_line = next(line for line in lines if line.startswith("Alice"))
    bob_line = next(line for line in lines if line.startswith("Bob"))

    # Both rows must have the same number of pipe-separated fields as the
    # header (4) -- Bob's blank Phone should be an empty field, not
    # missing entirely, so "Chennai" stays in the City column for both.
    assert alice_line.count("|") == bob_line.count("|") == 3
    assert bob_line.split("|")[-1].strip() == "Chennai"


def test_excel_extractor_raises_corrupt_file_error_on_bad_legacy_xls(tmp_path, mocker):
    import xlrd

    xls_path = tmp_path / "corrupt.xls"
    xls_path.write_bytes(b"garbage")

    mocker.patch("xlrd.open_workbook", side_effect=xlrd.XLRDError("Unsupported format"))

    with pytest.raises(CorruptFileError):
        ExcelExtractor().extract(_make_job(str(xls_path)))
