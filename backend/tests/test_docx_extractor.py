import pytest
from docx import Document

from app.db.models import ContentType, Job, JobStatus
from app.ingestion.docx_extractor import DocxExtractor
from app.ingestion.exceptions import CorruptFileError, NoExtractableTextError


def _make_job(path: str) -> Job:
    return Job(
        id="test-job",
        status=JobStatus.PENDING,
        content_type=ContentType.DOCX,
        source_type="file",
        source_reference="notes.docx",
        stored_file_path=path,
    )


def test_docx_extractor_reads_paragraph_text(tmp_path):
    doc_path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph("This is the first paragraph.")
    document.add_paragraph("This is the second paragraph.")
    document.save(str(doc_path))

    extracted = DocxExtractor().extract(_make_job(str(doc_path)))

    assert "first paragraph" in extracted.text
    assert "second paragraph" in extracted.text
    assert extracted.title == "notes"


def test_docx_extractor_raises_on_empty_document(tmp_path):
    doc_path = tmp_path / "empty.docx"
    Document().save(str(doc_path))

    with pytest.raises(NoExtractableTextError):
        DocxExtractor().extract(_make_job(str(doc_path)))


def test_docx_extractor_raises_on_corrupt_file(tmp_path):
    bad_path = tmp_path / "not_really.docx"
    bad_path.write_text("this is not a valid docx file")

    with pytest.raises(CorruptFileError):
        DocxExtractor().extract(_make_job(str(bad_path)))
