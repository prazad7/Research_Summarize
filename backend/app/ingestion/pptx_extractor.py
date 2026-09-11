from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.db.models import Job
from app.ingestion.base import BaseExtractor, ExtractedContent
from app.ingestion.exceptions import CorruptFileError, NoExtractableTextError


class PptxExtractor(BaseExtractor):
    def extract(self, job: Job) -> ExtractedContent:
        path = Path(job.stored_file_path)
        try:
            presentation = Presentation(str(path))
        except PackageNotFoundError as exc:
            raise CorruptFileError(original=exc) from exc

        slides_text: list[str] = []
        for i, slide in enumerate(presentation.slides, start=1):
            lines = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text.strip())
                if shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            lines.append(" | ".join(cells))
            if lines:
                slides_text.append(f"## Slide {i}\n" + "\n".join(lines))

        text = "\n\n".join(slides_text)
        if not text.strip():
            raise NoExtractableTextError()

        return ExtractedContent(
            text=text.strip(),
            title=path.stem,
            extra_metadata={"slide_count": len(presentation.slides)},
        )
