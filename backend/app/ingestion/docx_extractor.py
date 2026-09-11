from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from app.db.models import Job
from app.ingestion.base import BaseExtractor, ExtractedContent
from app.ingestion.exceptions import CorruptFileError, NoExtractableTextError


class DocxExtractor(BaseExtractor):
    def extract(self, job: Job) -> ExtractedContent:
        path = Path(job.stored_file_path)
        try:
            document = Document(str(path))
        except PackageNotFoundError as exc:
            raise CorruptFileError(original=exc) from exc

        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        table_lines = []
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    table_lines.append(" | ".join(cells))

        text = "\n\n".join(paragraphs + table_lines)
        if not text.strip():
            raise NoExtractableTextError()

        return ExtractedContent(text=text.strip(), title=path.stem)
