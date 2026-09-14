from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.db.models import Job
from app.ingestion.base import BaseExtractor, ExtractedContent
from app.ingestion.exceptions import CorruptFileError, NoExtractableTextError

_MAX_ROWS_PER_SHEET = 500  # keeps huge spreadsheets from blowing the LLM context


def _rows_to_sections(sheets: list[tuple[str, list[tuple], int]]) -> list[str]:
    """Shared by both the modern (.xlsx) and legacy (.xls) readers below --
    ``sheets`` is (sheet_title, rows, max_row) tuples, ``rows`` already
    values-only."""
    sections: list[str] = []
    for title, rows, max_row in sheets:
        lines = [f"## Sheet: {title}"]
        for i, row in enumerate(rows):
            if i >= _MAX_ROWS_PER_SHEET:
                lines.append(f"... ({max_row - _MAX_ROWS_PER_SHEET} more rows truncated)")
                break
            # A blank cell keeps its place as an empty field rather than
            # being dropped. Dropping it would shift every later value in
            # the row one column to the left -- e.g. if a middle column is
            # blank, the last column's value quietly ends up looking like
            # it belongs to the second-to-last column instead, corrupting
            # that row for anyone (human or LLM) reading it by position.
            cells = ["" if c is None else str(c) for c in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
        if len(lines) > 1:
            sections.append("\n".join(lines))
    return sections


def _extract_xlsx(path: Path) -> tuple[list[str], int]:
    """Modern OOXML workbooks (.xlsx/.xlsm), via openpyxl."""
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, zipfile.BadZipFile) as exc:
        # zipfile.BadZipFile is the one a plain `except OSError` misses --
        # .xlsx is a zip container, so anything that isn't a valid zip at
        # all (a legacy binary .xls mislabeled/mis-detected, a truncated
        # download, an unrelated file renamed to .xlsx) raises this instead.
        raise CorruptFileError(original=exc) from exc

    sheets = [
        (sheet.title, sheet.iter_rows(values_only=True), sheet.max_row)
        for sheet in workbook.worksheets
    ]
    return _rows_to_sections(sheets), len(workbook.worksheets)


def _extract_legacy_xls(path: Path) -> tuple[list[str], int]:
    """Legacy binary workbooks (.xls, pre-2007 format) -- openpyxl can only
    read the OOXML format, so these need a separate reader (xlrd, which
    dropped .xlsx support in 2.0 but still reads the old binary format)."""
    import xlrd

    try:
        workbook = xlrd.open_workbook(str(path))
    except (xlrd.XLRDError, OSError) as exc:
        raise CorruptFileError(original=exc) from exc

    sheets = [
        (sheet.name, (sheet.row_values(r) for r in range(sheet.nrows)), sheet.nrows)
        for sheet in workbook.sheets()
    ]
    return _rows_to_sections(sheets), workbook.nsheets


class ExcelExtractor(BaseExtractor):
    def extract(self, job: Job) -> ExtractedContent:
        path = Path(job.stored_file_path)

        if path.suffix.lower() == ".xls":
            sections, sheet_count = _extract_legacy_xls(path)
        else:
            sections, sheet_count = _extract_xlsx(path)

        text = "\n\n".join(sections)
        if not text.strip():
            raise NoExtractableTextError()

        return ExtractedContent(
            text=text.strip(),
            title=path.stem,
            extra_metadata={"sheet_count": sheet_count},
        )
